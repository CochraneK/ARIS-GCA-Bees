#!/usr/bin/env python3
"""Vendor the published Péron et al. 2019 S5 parameter table as UTF-8 CSV.

This converts the published S5 XLSX supporting file only. It does not contain
or reconstruct the restricted Species360 individual-level records.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


ARTICLE_DOI = "10.1371/journal.pbio.3000432"
SUPPLEMENT_DOI = "10.1371/journal.pbio.3000432.s005"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--csv", type=Path, required=True)
    ap.add_argument("--provenance", type=Path, required=True)
    ap.add_argument("--source-run-id", type=int, required=True)
    args = ap.parse_args()

    raw = args.xlsx.read_bytes()
    raw_sha = sha256_bytes(raw)

    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) != 97 or len(rows[0]) != 11:
        raise RuntimeError(
            f"Unexpected S5 shape: {len(rows)} rows x {len(rows[0])} columns"
        )

    args.csv.parent.mkdir(parents=True, exist_ok=True)
    with args.csv.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        for row in rows:
            writer.writerow(["" if value is None else value for value in row])

    csv_bytes = args.csv.read_bytes()
    provenance = {
        "source": "Péron et al. 2019, PLOS Biology, S5 Data",
        "article_doi": ARTICLE_DOI,
        "supplement_doi": SUPPLEMENT_DOI,
        "published_s5_description": (
            "Estimates and standard errors of the five parameters of the "
            "age-specific mortality curves for 96 species, with sample-size information."
        ),
        "retrieved_via_successful_github_actions_run": args.source_run_id,
        "original_xlsx_sha256": raw_sha,
        "derived_csv_sha256": sha256_bytes(csv_bytes),
        "rows_excluding_header": len(rows) - 1,
        "columns": len(rows[0]),
        "conversion": (
            "First XLSX worksheet read with openpyxl data_only=True; "
            "cell values preserved in UTF-8 CSV."
        ),
        "license": (
            "PLOS article is distributed under the Creative Commons Attribution "
            "License, permitting unrestricted use, distribution and reproduction "
            "with attribution."
        ),
        "restricted_data_guardrail": (
            "This vendored file is the published S5 parameter-estimate table. "
            "It is not the restricted Species360 individual-level longevity dataset."
        ),
    }
    args.provenance.write_text(
        json.dumps(provenance, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(provenance, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
