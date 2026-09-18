#!/usr/bin/env python3
"""ARIS4C007 Pilot 1: independent demographic-event benchmark.

Uses:
- official AnAge life-history traits for A1/A3 coordinate construction;
- Péron et al. 2019 S5 mortality-curve parameters as independent events.

The outcome is cross-species concentration of mapped ages for independently
estimated demographic events. It does not designate a globally "correct" age
mapping.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import statistics
import random
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from age_mappings import SpeciesTime, map_by_loglinear, map_by_relative_age
from build_pilot0 import read_anage, species_time


EST_RE = re.compile(r"^\s*([-+]?\d*\.?\d+(?:[eE][-+]?\d+)?)")


def estimate(value) -> float:
    if value is None:
        raise ValueError("missing estimate")
    if isinstance(value, (int, float)):
        return float(value)
    match = EST_RE.match(str(value))
    if not match:
        raise ValueError(f"cannot parse estimate: {value!r}")
    return float(match.group(1))


def quantile(values: list[float], p: float) -> float:
    xs = sorted(values)
    pos = (len(xs) - 1) * p
    lo = int(pos)
    hi = min(lo + 1, len(xs) - 1)
    frac = pos - lo
    return xs[lo] * (1 - frac) + xs[hi] * frac


def mad(values: list[float]) -> float:
    med = statistics.median(values)
    return statistics.median(abs(x - med) for x in values)


def dispersion(values: list[float]) -> dict[str, float]:
    return {
        "median_human_y": statistics.median(values),
        "mad_human_y": mad(values),
        "iqr_human_y": quantile(values, 0.75) - quantile(values, 0.25),
        "sd_human_y": statistics.stdev(values) if len(values) > 1 else 0.0,
        "min_human_y": min(values),
        "max_human_y": max(values),
    }


def bootstrap_mad_difference(
    a1: list[float],
    a3: list[float],
    *,
    seed: int,
    reps: int = 10000,
) -> dict[str, object]:
    """Paired species bootstrap for MAD(A3) - MAD(A1)."""
    if len(a1) != len(a3):
        raise ValueError("paired vectors must have equal length")
    rng = random.Random(seed)
    n = len(a1)
    diffs = []
    for _ in range(reps):
        idx = [rng.randrange(n) for _ in range(n)]
        b1 = [a1[i] for i in idx]
        b3 = [a3[i] for i in idx]
        diffs.append(mad(b3) - mad(b1))
    diffs.sort()
    return {
        "observed_mad_difference_A3_minus_A1_y": mad(a3) - mad(a1),
        "bootstrap_reps": reps,
        "bootstrap_95pct_interval_y": [
            quantile(diffs, 0.025),
            quantile(diffs, 0.975),
        ],
    }


def load_complete_anage(path: Path) -> dict[str, SpeciesTime]:
    out = {}
    for row in read_anage(path):
        if row.get("Class") != "Mammalia":
            continue
        s = species_time(row)
        if s is not None:
            out[s.name] = s
    return out


def load_peron(path: Path) -> list[dict[str, object]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(x) for x in rows[0]]
    return [dict(zip(header, row)) for row in rows[1:] if row[0]]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("anage", type=Path)
    ap.add_argument("peron_s5", type=Path)
    ap.add_argument("--out", type=Path, default=Path("data/pilot1_demography"))
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    anage = load_complete_anage(args.anage)
    human = anage["Homo sapiens"]
    peron = load_peron(args.peron_s5)

    output = []
    unmatched = []
    for row in peron:
        name = str(row["Species"]).strip()
        source = anage.get(name)
        if source is None:
            unmatched.append(name)
            continue

        juvenile_end = estimate(row["A"])
        prime_duration = estimate(row["Omegatilde"])
        senescence_onset = juvenile_end + prime_duration
        a10 = estimate(row["A10"])

        events = {
            "juvenile_stage_end_A": juvenile_end,
            "actuarial_senescence_onset_Omega": senescence_onset,
            "A10_10pct_survival": a10,
        }

        for event, age in events.items():
            a1 = map_by_relative_age(age, source, human)
            a3 = map_by_loglinear(age, source, human)
            output.append(
                {
                    "species": name,
                    "preferred_mortality_model": row["pref model"],
                    "event": event,
                    "source_event_age_y": age,
                    "A1_relative_lifespan_human_y": a1,
                    "A3_loglinear_human_y": a3,
                    "A3_minus_A1_y": a3 - a1,
                    "absolute_A1_A3_difference_y": abs(a3 - a1),
                }
            )

    csv_path = args.out / "demography_event_mapping.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(output[0]))
        writer.writeheader()
        writer.writerows(output)

    by_event = defaultdict(list)
    for row in output:
        by_event[row["event"]].append(row)

    event_summary = {}
    for event_index, (event, group) in enumerate(sorted(by_event.items())):
        a1 = [float(r["A1_relative_lifespan_human_y"]) for r in group]
        a3 = [float(r["A3_loglinear_human_y"]) for r in group]
        event_summary[event] = {
            "n_species": len(group),
            "A1_relative_lifespan": dispersion(a1),
            "A3_loglinear": dispersion(a3),
            "paired_bootstrap_mad_difference": bootstrap_mad_difference(
                a1,
                a3,
                seed=20260918 + event_index,
            ),
            "note": (
                "Lower cross-species dispersion is a necessary-style coherence "
                "diagnostic for a homologous event, not proof of biological truth."
            ),
        }

    summary = {
        "peron_species_rows": len(peron),
        "anage_complete_mammals_including_human": len(anage),
        "exact_overlap_species": len({r["species"] for r in output}),
        "unmatched_or_incomplete_species": unmatched,
        "event_summary": event_summary,
        "definitions": {
            "Omega": "A + Omegatilde, where Omegatilde is duration of the prime-age stage in the Péron table.",
            "A10": "predicted age at which 90% of the cohort is dead (10% survival).",
        },
        "guardrail": (
            "Péron demographic parameters are independent of the AnAge traits used "
            "to construct A1/A3, but population/captive conditions and event "
            "homology remain substantive limitations."
        ),
    }
    (args.out / "demography_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
