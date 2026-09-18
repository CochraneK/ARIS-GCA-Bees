#!/usr/bin/env python3
"""Acquire Januel et al. 2026 Translating Time supplementary sources.

Targets:
- Table S1: source observations/timepoints used to translate ages.
- Dataset 1: authors' R script that works with Table S1.

The script discovers official supplementary links from the peer-reviewed article
page/PMC record, validates bytes, converts Table S1 to UTF-8 CSV, extracts text
scripts, and writes provenance hashes. Raw downloaded binaries are temporary.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import shutil
import tarfile
import urllib.parse
import xml.etree.ElementTree as ET
import urllib.request
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ARTICLE_DOI = "10.1242/bio.062604"
PMCID = "PMC13382973"
PMC_VERSION = "1"
PMC_CLOUD_BASE = f"https://pmc-oa-opendata.s3.amazonaws.com/{PMCID}.{PMC_VERSION}/"
ARTICLE_URLS = [
    f"https://pmc.ncbi.nlm.nih.gov/articles/{PMCID}/",
    f"https://doi.org/{ARTICLE_DOI}",
]
TARGETS = {
    "table_s1": "biolopen-15-062604-TableS1.xlsx",
    "dataset1": "biolopen-15-062604-Dataset1.zip",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def get(url: str) -> tuple[bytes, str, str]:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": (
                "Mozilla/5.0 (compatible; ARIS4C007/1.0; "
                "+https://github.com/CochraneK/ARIS4C)"
            ),
            "Accept": "*/*",
        },
    )
    with urllib.request.urlopen(req, timeout=120) as response:
        return (
            response.read(),
            response.geturl(),
            response.headers.get("Content-Type", ""),
        )


def acquire_from_pmc_oa_package(tmp: Path) -> tuple[dict[str, Path], dict[str, object]]:
    """Acquire supplement bytes from the official PMC Open Access package."""
    oa_url = f"https://www.ncbi.nlm.nih.gov/pmc/utils/oa/oa.fcgi?id={PMCID}"
    body, final_url, content_type = get(oa_url)
    root = ET.fromstring(body)

    tgz_href = None
    for link in root.findall(".//link"):
        if link.attrib.get("format") == "tgz":
            tgz_href = link.attrib.get("href")
            break
    if not tgz_href:
        raise RuntimeError(
            "PMC OA API returned no tgz package link: "
            + body.decode("utf-8", errors="replace")[:500]
        )

    if tgz_href.startswith("ftp://ftp.ncbi.nlm.nih.gov/"):
        tgz_url = "https://ftp.ncbi.nlm.nih.gov/" + tgz_href.split(
            "ftp://ftp.ncbi.nlm.nih.gov/", 1
        )[1]
    else:
        tgz_url = tgz_href

    archive_bytes, archive_final_url, archive_type = get(tgz_url)
    tgz_path = tmp / f"{PMCID}.tar.gz"
    tgz_path.write_bytes(archive_bytes)

    found: dict[str, Path] = {}
    members_seen = []
    with tarfile.open(tgz_path, "r:gz") as tf:
        for member in tf.getmembers():
            if not member.isfile():
                continue
            basename = Path(member.name).name
            members_seen.append(member.name)
            for key, filename in TARGETS.items():
                if basename == filename:
                    extracted = tf.extractfile(member)
                    if extracted is None:
                        continue
                    dest = tmp / filename
                    dest.write_bytes(extracted.read())
                    found[key] = dest

    missing = [key for key in TARGETS if key not in found]
    if missing:
        raise RuntimeError(
            f"PMC OA package missing targets {missing}; "
            f"sample members={members_seen[:40]}"
        )

    meta = {
        "oa_api_url": oa_url,
        "oa_api_final_url": final_url,
        "oa_api_content_type": content_type,
        "oa_package_url": tgz_url,
        "oa_package_final_url": archive_final_url,
        "oa_package_content_type": archive_type,
        "oa_package_bytes": len(archive_bytes),
        "oa_package_sha256": sha256(tgz_path),
    }
    return found, meta


def discover_links() -> tuple[dict[str, str], list[dict[str, object]]]:
    found: dict[str, str] = {}
    attempts: list[dict[str, object]] = []

    for article_url in ARTICLE_URLS:
        try:
            body, final_url, content_type = get(article_url)
            text = body.decode("utf-8", errors="replace")
            attempts.append(
                {
                    "requested_url": article_url,
                    "final_url": final_url,
                    "bytes": len(body),
                    "content_type": content_type,
                    "success": True,
                }
            )
        except Exception as exc:
            attempts.append(
                {
                    "requested_url": article_url,
                    "success": False,
                    "error": f"{type(exc).__name__}: {exc}",
                }
            )
            continue

        # href extraction is deliberately liberal because PMC and journal pages
        # may encode ampersands or wrap supplementary links differently.
        hrefs = re.findall(
            r"""href\s*=\s*["']([^"']+)["']""",
            text,
            flags=re.IGNORECASE,
        )
        hrefs = [html.unescape(x) for x in hrefs]

        for key, filename in TARGETS.items():
            for href in hrefs:
                if filename.lower() in href.lower():
                    found[key] = urllib.parse.urljoin(final_url, href)
                    break

        if len(found) == len(TARGETS):
            break

    # Known PMC relative-path fallbacks. These are attempted only if automatic
    # discovery fails and are still byte-validated downstream.
    for key, filename in TARGETS.items():
        if key not in found:
            found[key] = (
                f"https://pmc.ncbi.nlm.nih.gov/articles/{PMCID}/bin/{filename}"
            )

    return found, attempts


def download_validated(urls: list[str], dest: Path, kind: str) -> dict[str, object]:
    errors = []
    for url in urls:
        try:
            body, final_url, content_type = get(url)
            dest.write_bytes(body)

            valid_zip = zipfile.is_zipfile(dest)
            if kind == "xlsx":
                if not valid_zip:
                    raise ValueError("downloaded Table S1 is not an XLSX/ZIP container")
                with zipfile.ZipFile(dest) as zf:
                    names = zf.namelist()
                    if "[Content_Types].xml" not in names:
                        raise ValueError("Table S1 XLSX validation failed")
            elif kind == "zip":
                if not valid_zip:
                    raise ValueError("downloaded Dataset 1 is not a ZIP")
            else:
                raise ValueError(kind)

            return {
                "requested_url": url,
                "final_url": final_url,
                "content_type": content_type,
                "bytes": len(body),
                "sha256": sha256(dest),
            }
        except Exception as exc:
            errors.append(f"{url}: {type(exc).__name__}: {exc}")
            dest.unlink(missing_ok=True)

    raise RuntimeError("All supplementary download attempts failed:\n" + "\n".join(errors))


def convert_table_s1(xlsx: Path, out_dir: Path) -> list[dict[str, object]]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    outputs = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        safe = re.sub(r"[^A-Za-z0-9._-]+", "_", sheet_name).strip("_") or "sheet"
        csv_path = out_dir / f"TableS1__{safe}.csv"

        row_count = 0
        max_columns = 0
        with csv_path.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh, lineterminator="\n")
            for row in ws.iter_rows(values_only=True):
                values = ["" if x is None else x for x in row]
                writer.writerow(values)
                row_count += 1
                max_columns = max(max_columns, len(values))

        outputs.append(
            {
                "sheet": sheet_name,
                "file": csv_path.name,
                "rows_including_header": row_count,
                "max_columns": max_columns,
                "sha256": sha256(csv_path),
            }
        )
    return outputs


def extract_dataset1(zip_path: Path, out_dir: Path) -> list[dict[str, object]]:
    extracted = []
    scripts_dir = out_dir / "Dataset1"
    scripts_dir.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(zip_path) as zf:
        members = [x for x in zf.namelist() if not x.endswith("/")]
        for member in members:
            raw = zf.read(member)
            basename = Path(member).name
            safe = re.sub(r"[^A-Za-z0-9._-]+", "_", basename)
            if not safe:
                continue
            path = scripts_dir / safe
            path.write_bytes(raw)
            extracted.append(
                {
                    "source_member": member,
                    "file": f"Dataset1/{safe}",
                    "bytes": len(raw),
                    "sha256": sha256(path),
                }
            )

    if not extracted:
        raise RuntimeError("Dataset 1 ZIP contained no files")
    return extracted


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()

    args.out.mkdir(parents=True, exist_ok=True)
    tmp = args.out / "_tmp"
    shutil.rmtree(tmp, ignore_errors=True)
    tmp.mkdir(parents=True)

    discovery_attempts = []
    oa_package_meta = None

    # PMC changed its Article Dataset Distribution Services in August 2026.
    # Try the current AWS Cloud Service first; retain legacy fallbacks only for
    # historical reproducibility.
    cloud_table = urllib.parse.urljoin(PMC_CLOUD_BASE, TARGETS["table_s1"])
    cloud_dataset = urllib.parse.urljoin(PMC_CLOUD_BASE, TARGETS["dataset1"])
    table_path = tmp / TARGETS["table_s1"]
    dataset_path = tmp / TARGETS["dataset1"]

    try:
        table_meta = download_validated([cloud_table], table_path, "xlsx")
        dataset_meta = download_validated([cloud_dataset], dataset_path, "zip")
        table_meta["acquisition"] = "PMC AWS Cloud Service"
        dataset_meta["acquisition"] = "PMC AWS Cloud Service"
        oa_package_meta = {
            "service": "PMC AWS Cloud Service",
            "base_url": PMC_CLOUD_BASE,
            "article_version": f"{PMCID}.{PMC_VERSION}",
        }
    except Exception as cloud_exc:
        discovery_attempts.append(
            {
                "source": "PMC AWS Cloud Service",
                "success": False,
                "error": f"{type(cloud_exc).__name__}: {cloud_exc}",
            }
        )

        try:
            oa_files, oa_package_meta = acquire_from_pmc_oa_package(tmp)
            table_path = oa_files["table_s1"]
            dataset_path = oa_files["dataset1"]

            if not zipfile.is_zipfile(table_path):
                raise RuntimeError("PMC OA Table S1 is not a valid XLSX container")
            if not zipfile.is_zipfile(dataset_path):
                raise RuntimeError("PMC OA Dataset 1 is not a valid ZIP container")

            table_meta = {
            "acquisition": "PMC Open Access package",
            "file": table_path.name,
            "bytes": table_path.stat().st_size,
            "sha256": sha256(table_path),
        }
            dataset_meta = {
                "acquisition": "PMC Open Access package",
            "file": dataset_path.name,
            "bytes": dataset_path.stat().st_size,
            "sha256": sha256(dataset_path),
        }
        except Exception as oa_exc:
            discovery_attempts.append(
            {
                    "source": "PMC Open Access package",
                    "success": False,
                    "error": f"{type(oa_exc).__name__}: {oa_exc}",
                }
            )

            links, html_attempts = discover_links()
            discovery_attempts.extend(html_attempts)

            table_path = tmp / TARGETS["table_s1"]
            dataset_path = tmp / TARGETS["dataset1"]

            table_urls = [
            urllib.parse.urljoin(PMC_CLOUD_BASE, TARGETS["table_s1"]),
            links["table_s1"],
            f"https://pmc.ncbi.nlm.nih.gov/articles/{PMCID}/bin/{TARGETS['table_s1']}",
            f"https://pmc.ncbi.nlm.nih.gov/articles/instance/{PMCID.replace('PMC','')}/bin/{TARGETS['table_s1']}",
        ]
            dataset_urls = [
            urllib.parse.urljoin(PMC_CLOUD_BASE, TARGETS["dataset1"]),
            links["dataset1"],
            f"https://pmc.ncbi.nlm.nih.gov/articles/{PMCID}/bin/{TARGETS['dataset1']}",
            f"https://pmc.ncbi.nlm.nih.gov/articles/instance/{PMCID.replace('PMC','')}/bin/{TARGETS['dataset1']}",
        ]

            table_meta = download_validated(
                list(dict.fromkeys(table_urls)),
                table_path,
                "xlsx",
            )
            dataset_meta = download_validated(
                list(dict.fromkeys(dataset_urls)),
                dataset_path,
                "zip",
            )

    table_outputs = convert_table_s1(table_path, args.out)
    dataset_outputs = extract_dataset1(dataset_path, args.out)

    provenance = {
        "study": "Januel et al. 2026, Biology Open",
        "title": (
            "Cat brains age like humans: translating time shows pet cats live "
            "to be natural models for human aging"
        ),
        "article_doi": ARTICLE_DOI,
        "pmcid": PMCID,
        "license": "CC BY 4.0",
        "supplementary_targets": {
            "TableS1": {
                **table_meta,
                "description": (
                    "Timepoints and observations in cats, humans, mice and "
                    "chimpanzees used to translate ages across species."
                ),
                "derived_text_files": table_outputs,
            },
            "Dataset1": {
                **dataset_meta,
                "description": (
                    "Authors' script bundle used with Table S1 to translate ages "
                    "across species."
                ),
                "extracted_files": dataset_outputs,
            },
        },
        "pmc_oa_package": oa_package_meta,
        "link_discovery_attempts": discovery_attempts,
        "guardrail": (
            "These are source supplementary materials. ARIS4C007 will preserve "
            "the authors' model as a benchmark and separately evaluate held-out "
            "events to avoid treating training fit as independent validation."
        ),
    }

    (args.out / "provenance.json").write_text(
        json.dumps(provenance, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    shutil.rmtree(tmp)
    print(json.dumps(provenance, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
